"""Test OCR fallback trong ingestion.py — dùng tesseract THẬT (không mock),
vì đây đúng là phần dễ âm thầm hỏng nhất (đổi ngưỡng, đổi API pytesseract,
thiếu binary tesseract-ocr trên máy CI/production) mà một bản mock sẽ
không bao giờ bắt được. CI workflow (ci.yml) cài tesseract-ocr +
tesseract-ocr-vie qua apt trước khi chạy bộ test này.
"""
import docx
import fitz
import pytesseract
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches

from app.rag.ingestion import extract_text


def _make_text_image(path: str, text: str) -> None:
    image = Image.new("RGB", (600, 100), color="white")
    draw = ImageDraw.Draw(image)
    draw.text((10, 30), text, fill="black")
    image.save(path)


def test_pdf_with_real_text_layer_skips_ocr(tmp_path):
    pdf_path = tmp_path / "text.pdf"
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 50), "Dinh luat 2 Newton: F = m * a")
    doc.save(str(pdf_path))
    doc.close()

    result = extract_text(str(pdf_path), "application/pdf")
    assert "Newton" in result


def test_pdf_with_no_text_layer_falls_back_to_ocr(tmp_path):
    image_path = tmp_path / "source.png"
    _make_text_image(str(image_path), "Cong thuc hoa hoc NaCl")

    pdf_path = tmp_path / "scanned.pdf"
    doc = fitz.open()
    page = doc.new_page()
    page.insert_image(fitz.Rect(50, 50, 550, 150), filename=str(image_path))
    doc.save(str(pdf_path))
    doc.close()

    result = extract_text(str(pdf_path), "application/pdf")
    # Không đòi hỏi OCR chính xác 100% (phụ thuộc font/chất lượng ảnh —
    # font bitmap mặc định của PIL ở kích thước nhỏ có thể đọc nhầm vài
    # ký tự, ví dụ "hoa" -> "hea") — chỉ cần xác nhận pipeline THỰC SỰ
    # chạy OCR và trả về nội dung liên quan, thay vì "final_text rỗng"
    # như trước khi có fallback. "NaCl" là chuỗi viết hoa/thường rõ ràng,
    # ít bị OCR đọc nhầm hơn các từ tiếng Việt không dấu.
    assert result.strip() != ""
    assert "nacl" in result.lower()


def test_docx_paragraph_and_embedded_image_both_extracted(tmp_path):
    image_path = tmp_path / "diagram.png"
    _make_text_image(str(image_path), "Luc ma sat")

    docx_path = tmp_path / "notes.docx"
    d = docx.Document()
    d.add_paragraph("Ly thuyet co hoc co ban")
    d.add_picture(str(image_path))
    d.save(str(docx_path))

    mime = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    result = extract_text(str(docx_path), mime)

    assert "Ly thuyet co hoc co ban" in result
    assert result.strip() != "Ly thuyet co hoc co ban"  # OCR text cũng phải có mặt


def test_pptx_slide_text_and_embedded_image_both_extracted(tmp_path):
    image_path = tmp_path / "diagram.png"
    _make_text_image(str(image_path), "Dinh luat bao toan")

    pptx_path = tmp_path / "slides.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    textbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    textbox.text_frame.text = "Vat ly dai cuong"
    slide.shapes.add_picture(str(image_path), Inches(1), Inches(2.5), width=Inches(5))
    prs.save(str(pptx_path))

    mime = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
    result = extract_text(str(pptx_path), mime)

    assert "Vat ly dai cuong" in result
    assert result.strip() != "Vat ly dai cuong"  # OCR text cũng phải có mặt


def test_ocr_failure_degrades_gracefully_without_crashing(tmp_path, monkeypatch):
    """Nếu binary tesseract-ocr không có trên máy (thiếu cài đặt hệ điều
    hành), pipeline không được crash toàn bộ document — chỉ trả về ít/không
    có text OCR, để worker/tasks.py xử lý như "không trích xuất được nội
    dung" (status=failed) thay vì lỗi 500 khó hiểu.
    """
    monkeypatch.setattr(pytesseract.pytesseract, "tesseract_cmd", "/nonexistent/tesseract")

    image_path = tmp_path / "source.png"
    _make_text_image(str(image_path), "Bat ky noi dung gi")

    pdf_path = tmp_path / "scanned.pdf"
    doc = fitz.open()
    page = doc.new_page()
    page.insert_image(fitz.Rect(50, 50, 550, 150), filename=str(image_path))
    doc.save(str(pdf_path))
    doc.close()

    # Không raise exception — đây là điều quan trọng nhất cần test.
    result = extract_text(str(pdf_path), "application/pdf")
    assert result == ""


def test_markdown_is_extracted_and_chunked(tmp_path):
    """.md is an allowed upload type; it must not fail the ingestion worker."""
    from app.rag.ingestion import extract_and_chunk

    md_path = tmp_path / "notes.md"
    md_path.write_text(
        "# OOP\nClasses are blueprints for objects in Java programs.\n"
        "```python\n# not a heading\n```\n"
        "## Inheritance\nA subclass extends a superclass and reuses behaviour.\n",
        encoding="utf-8",
    )
    # Browsers may send .md as text/markdown or application/octet-stream.
    for mime in ("text/markdown", "application/octet-stream"):
        chunks = extract_and_chunk(str(md_path), mime)
        assert chunks
        assert {c.section_title for c in chunks} >= {"OOP", "Inheritance"}
