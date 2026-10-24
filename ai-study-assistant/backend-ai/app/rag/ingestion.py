"""Document ingestion — text extraction + structured chunking.

Returns typed ``ChunkData`` objects (from chunking.py) that carry source
metadata (page number, slide number/title, section heading) so the vector
store can persist them and the chat endpoint can render precise citations.

Also supports a *summarization phase*: after extracting all chunks for a
document, a short extractive summary is generated per group of segments
(PDF page or PPTX slide) and stored as an extra ``ChunkData`` with
``extra["is_summary"] = True``.  This summary chunk lets the retriever find
high-level answers without pulling in many noisy detail chunks.
"""
from __future__ import annotations

import io
import logging
import re
from typing import Any, Dict, List

import docx
import fitz
import pytesseract
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from app.rag.chunking import (
    ChunkData,
    chunk_docx_sections,
    chunk_pdf_pages,
    chunk_pptx_slides,
    chunk_text,
)

logger = logging.getLogger(__name__)

# If text extracted from a page/block is shorter than this, try OCR instead.
OCR_FALLBACK_THRESHOLD_CHARS = 20
OCR_LANGUAGES = "vie+eng"

# Minimum character count for a summary segment to be worth generating.
_SUMMARY_MIN_CHARS = 200


# OCR helper
def _ocr_image_bytes(image_bytes: bytes) -> str:
    """Run basic OCR on raw image bytes; return '' on any failure."""
    try:
        image = Image.open(io.BytesIO(image_bytes))
        return pytesseract.image_to_string(image, lang=OCR_LANGUAGES)
    except pytesseract.TesseractNotFoundError:
        logger.warning(
            "tesseract-ocr binary not found — skipping OCR. "
            "Install tesseract-ocr (+ tesseract-ocr-vie) at OS level."
        )
        return ""
    except Exception as exc:  # noqa: BLE001
        logger.warning(f"OCR failed on 1 image: {exc}")
        return ""

# Per-format structured extraction
def _extract_pdf_pages(file_path: str) -> List[str]:
    """Return one text string per page (may be OCR-derived for scanned pages)."""
    pages: List[str] = []
    with fitz.open(file_path) as doc:
        for page in doc:
            text = page.get_text("text")
            if text and len(text.strip()) >= OCR_FALLBACK_THRESHOLD_CHARS:
                pages.append(text)
                continue
            # Likely a scanned page — render and OCR.
            pixmap = page.get_pixmap(dpi=200)
            ocr_text = _ocr_image_bytes(pixmap.tobytes("png"))
            if ocr_text.strip():
                pages.append(ocr_text)
            elif text:
                pages.append(text)  # keep whatever little text exists
            else:
                pages.append("")   # blank page placeholder keeps index stable
    return pages


def _extract_docx_sections(file_path: str) -> List[Dict[str, Any]]:
    """Return a list of section dicts, one per heading (or the whole doc if no headings)."""
    document = docx.Document(file_path)

    sections: List[Dict[str, Any]] = []
    current_heading: str | None = None
    current_paragraphs: List[str] = []

    def _flush():
        text = "\n".join(current_paragraphs).strip()
        if text:
            sections.append({"heading": current_heading, "text": text})

    for para in document.paragraphs:
        if para.style.name.startswith("Heading"):
            _flush()
            current_heading = para.text.strip() or None
            current_paragraphs = []
        elif para.text.strip():
            current_paragraphs.append(para.text)

    _flush()

    # OCR embedded images
    for shape in document.inline_shapes:
        try:
            blip = shape._inline.graphic.graphicData.pic.blipFill.blip
            image_part = document.part.related_parts[blip.embed]
            ocr_text = _ocr_image_bytes(image_part.blob)
            if ocr_text.strip():
                sections.append({"heading": None, "text": ocr_text})
        except Exception as exc:  # noqa: BLE001
            logger.warning(f"Cannot read 1 embedded image in DOCX: {exc}")

    return sections if sections else [{"heading": None, "text": ""}]


def _is_markdown(file_path: str, mime_type: str) -> bool:
    """Markdown is accepted by the upload endpoint; browsers often send it as
    text/markdown or application/octet-stream, so also check the extension."""
    return file_path.lower().endswith((".md", ".markdown")) or "markdown" in mime_type.lower()


def _extract_md_sections(file_path: str) -> List[Dict[str, Any]]:
    """Return section dicts (same shape as DOCX), split on ``#`` headings.

    Lines inside fenced code blocks are never treated as headings.
    """
    with open(file_path, "r", encoding="utf-8-sig", errors="replace") as fh:
        lines = fh.read().splitlines()

    sections: List[Dict[str, Any]] = []
    current_heading: str | None = None
    current_lines: List[str] = []
    in_fence = False

    def _flush():
        text = "\n".join(current_lines).strip()
        if text:
            sections.append({"heading": current_heading, "text": text})

    for line in lines:
        if line.lstrip().startswith(("```", "~~~")):
            in_fence = not in_fence
            current_lines.append(line)
            continue
        m = None if in_fence else re.match(r"^#{1,6}\s+(.*?)\s*#*\s*$", line)
        if m:
            _flush()
            current_heading = m.group(1).strip() or None
            current_lines = []
        else:
            current_lines.append(line)

    _flush()
    return sections if sections else [{"heading": None, "text": ""}]


def _extract_pptx_slides(file_path: str) -> List[Dict[str, Any]]:
    """Return a list of slide dicts with slide_number, title, text."""
    prs = Presentation(file_path)
    slides: List[Dict[str, Any]] = []

    for slide_num, slide in enumerate(prs.slides, start=1):
        title_text: str | None = None
        body_parts: List[str] = []

        for shape in slide.shapes:
            # Extract slide title from the title placeholder
            if shape.has_text_frame and shape.name.lower().startswith("title"):
                t = shape.text.strip()
                if t:
                    title_text = t
                    continue
            if hasattr(shape, "text") and shape.text.strip():
                body_parts.append(shape.text.strip())
            elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                try:
                    ocr_text = _ocr_image_bytes(shape.image.blob)
                    if ocr_text.strip():
                        body_parts.append(ocr_text)
                except Exception as exc:  # noqa: BLE001
                    logger.warning(f"Cannot read 1 image in PPTX slide {slide_num}: {exc}")

        slides.append(
            {
                "slide_number": slide_num,
                "title": title_text,
                "text": "\n".join(body_parts),
            }
        )

    return slides


# Summarization helper (extractive, runs at ingestion time)
def _make_summary_chunk(
    text: str,
    chunk_index: int,
    page_number: int | None = None,
    slide_number: int | None = None,
    slide_title: str | None = None,
    section_title: str | None = None,
) -> ChunkData | None:
    """Create a short extractive summary chunk for a segment of text.

    Strategy: take the first 3 sentences of the segment as a proxy summary.
    This is a purely local operation — no LLM call — so it costs zero latency
    at ingestion time and zero runtime tokens at query time.
    """
    text = text.strip()
    if len(text) < _SUMMARY_MIN_CHARS:
        return None  # too short to summarise

    # Split on sentence boundaries (naive but language-agnostic).
    import re
    sentences = re.split(r"(?<=[.!?])\s+", text)
    summary_text = " ".join(sentences[:3]).strip()
    if not summary_text or len(summary_text) < 20:
        return None

    return ChunkData(
        content=summary_text,
        page_number=page_number,
        slide_number=slide_number,
        slide_title=slide_title,
        section_title=section_title,
        chunk_index=chunk_index,
        extra={"is_summary": True},
    )


# Public API
def extract_and_chunk(file_path: str, mime_type: str) -> List[ChunkData]:
    """Extract text from a document and return structured ChunkData objects.

    This is the preferred entry point for the ingestion worker.  It returns
    metadata-rich chunks AND optional per-segment summary chunks appended at
    the end.
    """
    try:
        if "pdf" in mime_type.lower() or file_path.lower().endswith(".pdf"):
            pages = _extract_pdf_pages(file_path)
            detail_chunks = chunk_pdf_pages(pages)
            # Build per-page summaries
            summary_chunks: List[ChunkData] = []
            idx_offset = len(detail_chunks)
            for page_num, page_text in enumerate(pages, start=1):
                sc = _make_summary_chunk(
                    page_text,
                    chunk_index=idx_offset,
                    page_number=page_num,
                )
                if sc:
                    summary_chunks.append(sc)
                    idx_offset += 1
            return detail_chunks + summary_chunks

        elif (
            _is_markdown(file_path, mime_type)
            or "word" in mime_type.lower()
            or file_path.lower().endswith(".docx")
        ):
            if _is_markdown(file_path, mime_type):
                sections = _extract_md_sections(file_path)
            else:
                sections = _extract_docx_sections(file_path)
            detail_chunks = chunk_docx_sections(sections)
            summary_chunks = []
            idx_offset = len(detail_chunks)
            for section in sections:
                sc = _make_summary_chunk(
                    section.get("text", ""),
                    chunk_index=idx_offset,
                    section_title=section.get("heading"),
                )
                if sc:
                    summary_chunks.append(sc)
                    idx_offset += 1
            return detail_chunks + summary_chunks

        elif (
            "presentation" in mime_type.lower()
            or file_path.lower().endswith(".pptx")
            or file_path.lower().endswith(".ppt")
        ):
            slides = _extract_pptx_slides(file_path)
            detail_chunks = chunk_pptx_slides(slides)
            summary_chunks = []
            idx_offset = len(detail_chunks)
            for slide in slides:
                sc = _make_summary_chunk(
                    slide.get("text", ""),
                    chunk_index=idx_offset,
                    slide_number=slide.get("slide_number"),
                    slide_title=slide.get("title"),
                )
                if sc:
                    summary_chunks.append(sc)
                    idx_offset += 1
            return detail_chunks + summary_chunks

        else:
            raise ValueError(f"Unsupported file format: {mime_type}")

    except Exception as e:
        logger.error(f"Critical error extracting text from {file_path}: {e}")
        raise


# Legacy flat-text API (kept for backward compatibility with existing tests)
def extract_text(file_path: str, mime_type: str) -> str:
    """Return all extracted text as a single string (no structural metadata).

    Prefer ``extract_and_chunk`` for new code.
    """
    try:
        if "pdf" in mime_type.lower() or file_path.lower().endswith(".pdf"):
            pages = _extract_pdf_pages(file_path)
            return "\n".join(pages).strip()
        elif _is_markdown(file_path, mime_type):
            sections = _extract_md_sections(file_path)
            return "\n".join(s.get("text", "") for s in sections).strip()
        elif "word" in mime_type.lower() or file_path.lower().endswith(".docx"):
            sections = _extract_docx_sections(file_path)
            return "\n".join(s.get("text", "") for s in sections).strip()
        elif (
            "presentation" in mime_type.lower()
            or file_path.lower().endswith(".pptx")
            or file_path.lower().endswith(".ppt")
        ):
            slides = _extract_pptx_slides(file_path)
            return "\n".join(s.get("text", "") for s in slides).strip()
        else:
            raise ValueError(f"Unsupported file format: {mime_type}")
    except Exception as e:
        logger.error(f"Critical error extracting text from {file_path}: {e}")
        raise
